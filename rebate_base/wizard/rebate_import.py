"""T39 – import running agreements from an Excel file (one row per tier)."""
import base64
import io
from collections import OrderedDict

import openpyxl

from odoo import fields, models
from odoo.exceptions import UserError
from odoo.fields import Command

from ..models import selections as sel

COLUMNS = [
    ('agreement_ref', 'Mã thỏa thuận (gom dòng)', 'TT-NPP-A-2026'),
    ('title', 'Tên chương trình', 'Thưởng doanh số NPP A 2026'),
    ('direction', 'Chiều (sale/purchase)', 'sale'),
    ('partners', 'Đối tác (MST hoặc tên, cách nhau ;)', '0101234567'),
    ('contract_ref', 'Số hợp đồng', 'HĐ 12/2026'),
    ('date_from', 'Từ ngày (dd/mm/yyyy)', '01/01/2026'),
    ('date_to', 'Đến ngày (dd/mm/yyyy)', '31/12/2026'),
    ('period_type', 'Kỳ tính', 'quarter'),
    ('settlement_doc', 'Chứng từ quyết toán', 'adjust_invoice'),
    ('payment_method', 'Hình thức chi trả', 'offset'),
    ('rule_name', 'Tên quy tắc', 'Tất cả sản phẩm'),
    ('product_scope', 'Phạm vi (all/category/product)', 'all'),
    ('scope_values', 'Nhóm SP / mã SP (cách nhau ;)', ''),
    ('basis', 'Cơ sở tính', 'amount'),
    ('tier_mode', 'Kiểu bậc (retro/progressive)', 'retro'),
    ('value_type', 'Giá trị (percent/amount/per_unit)', 'percent'),
    ('tier_from', 'Bậc từ', 1_000_000_000),
    ('tier_to', 'Bậc đến (0 = không giới hạn)', 2_000_000_000),
    ('tier_value', 'Giá trị bậc', 1),
]
SELECTION_FIELDS = {
    'direction': sel.DIRECTION, 'period_type': sel.PERIOD_TYPE, 'settlement_doc': sel.SETTLEMENT_DOC,
    'payment_method': sel.PAYMENT_METHOD, 'basis': sel.BASIS, 'tier_mode': sel.TIER_MODE,
    'value_type': sel.VALUE_TYPE,
}


class RebateImport(models.TransientModel):
    _name = 'rebate.import'
    _description = 'Nhập thỏa thuận chiết khấu từ Excel'

    file_data = fields.Binary('File Excel')
    file_name = fields.Char()
    template_data = fields.Binary('File mẫu', readonly=True)
    template_name = fields.Char(default='mau_nhap_thoa_thuan_ck.xlsx')
    approve = fields.Boolean('Duyệt ngay và dựng lũy kế từ hóa đơn đã ghi sổ', default=True)
    result = fields.Text('Kết quả', readonly=True)

    # ------------------------------------------------------------------
    def _template_bytes(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'ThoaThuan'
        ws.append([c[0] for c in COLUMNS])
        ws.append([c[1] for c in COLUMNS])
        sample = [c[2] for c in COLUMNS]
        ws.append(sample)
        ws.append(sample[:16] + [2_000_000_000, 0, 2])
        guide = wb.create_sheet('HuongDan')
        guide.append(['Cột', 'Giá trị hợp lệ'])
        for key, values in SELECTION_FIELDS.items():
            guide.append([key, ', '.join('%s (%s)' % (v, label) for v, label in values)])
        guide.append(['', 'Dòng 1 là mã cột (không sửa), dòng 2 là diễn giải, dữ liệu bắt đầu từ dòng 3.'])
        guide.append(['', 'Mỗi dòng là một bậc. Các dòng cùng "agreement_ref" + "rule_name" thuộc cùng một quy tắc.'])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def action_download_template(self):
        self.ensure_one()
        self.template_data = base64.b64encode(self._template_bytes())
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/rebate.import/%s/template_data/%s?download=true' % (self.id, self.template_name),
            'target': 'self',
        }

    # ------------------------------------------------------------------
    @staticmethod
    def _to_date(value):
        if not value:
            return False
        if hasattr(value, 'date'):
            return value.date()
        text = str(value).strip()
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                from datetime import datetime
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        raise UserError(f'Ngày không hợp lệ: {text}')

    def _find_partners(self, text):
        Partner = self.env['res.partner']
        partners = Partner
        for token in [t.strip() for t in str(text or '').split(';') if t.strip()]:
            partner = Partner.search([('vat', '=', token), ('is_company', '=', True)], limit=1) \
                or Partner.search([('vat', '=', token)], limit=1) \
                or Partner.search([('name', '=ilike', token)], limit=1)
            if not partner:
                raise UserError(f'Không tìm thấy đối tác "{token}".')
            partners |= partner
        return partners

    def _scope_vals(self, scope, text):
        tokens = [t.strip() for t in str(text or '').split(';') if t.strip()]
        if scope == 'category':
            categs = self.env['product.category']
            for token in tokens:
                categ = categs.search(['|', ('complete_name', '=ilike', token), ('name', '=ilike', token)], limit=1)
                if not categ:
                    raise UserError(f'Không tìm thấy nhóm sản phẩm "{token}".')
                categs |= categ
            return {'product_categ_ids': [Command.set(categs.ids)]}
        if scope == 'product':
            products = self.env['product.product']
            for token in tokens:
                product = products.search(['|', ('default_code', '=', token), ('name', '=ilike', token)], limit=1)
                if not product:
                    raise UserError(f'Không tìm thấy sản phẩm "{token}".')
                products |= product
            return {'product_ids': [Command.set(products.ids)]}
        return {}

    def _check_selection(self, key, value, row_no):
        allowed = [v for v, _l in SELECTION_FIELDS[key]]
        if value not in allowed:
            raise UserError(f'Dòng {row_no}: "{key}" = "{value}" không hợp lệ ({", ".join(allowed)}).')
        return value

    def _parse(self, data):
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise UserError('File trống.')
        header = [str(h).strip() if h else '' for h in rows[0]]
        missing = [c[0] for c in COLUMNS if c[0] not in header]
        if missing:
            raise UserError('Thiếu cột: %s' % ', '.join(missing))
        idx = {name: header.index(name) for name, _l, _s in COLUMNS}
        agreements = OrderedDict()
        for row_no, row in enumerate(rows[2:], start=3):
            if not row or not any(row):
                continue
            get = lambda k, r=row: r[idx[k]] if idx[k] < len(r) else None  # noqa: E731
            ref = str(get('agreement_ref') or '').strip()
            if not ref:
                raise UserError(f'Dòng {row_no}: thiếu mã thỏa thuận.')
            agr = agreements.setdefault(ref, {'row': row_no, 'rules': OrderedDict(), 'vals': {
                'title': get('title') or ref,
                'direction': self._check_selection('direction', str(get('direction') or 'sale').strip(), row_no),
                'partners': get('partners'),
                'contract_ref': get('contract_ref') or ref,
                'date_from': self._to_date(get('date_from')),
                'date_to': self._to_date(get('date_to')),
                'period_type': self._check_selection('period_type', str(get('period_type') or 'quarter').strip(),
                                                     row_no),
                'settlement_doc': self._check_selection(
                    'settlement_doc', str(get('settlement_doc') or 'adjust_invoice').strip(), row_no),
                'payment_method': self._check_selection(
                    'payment_method', str(get('payment_method') or 'offset').strip(), row_no),
            }})
            rule_name = str(get('rule_name') or 'Quy tắc').strip()
            scope = str(get('product_scope') or 'all').strip()
            rule = agr['rules'].setdefault(rule_name, {'tiers': [], 'vals': {
                'name': rule_name,
                'product_scope': scope,
                'basis': self._check_selection('basis', str(get('basis') or 'amount').strip(), row_no),
                'tier_mode': self._check_selection('tier_mode', str(get('tier_mode') or 'retro').strip(), row_no),
                'value_type': self._check_selection('value_type', str(get('value_type') or 'percent').strip(),
                                                    row_no),
                **self._scope_vals(scope, get('scope_values')),
            }})
            rule['tiers'].append({
                'from_value': float(get('tier_from') or 0),
                'to_value': float(get('tier_to') or 0),
                'value': float(get('tier_value') or 0),
            })
        return agreements

    def action_import(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError('Hãy chọn file Excel.')
        parsed = self._parse(base64.b64decode(self.file_data))
        Agreement = self.env['rebate.agreement']
        created = Agreement
        for ref, agr in parsed.items():
            vals = dict(agr['vals'])
            partners = self._find_partners(vals.pop('partners'))
            vals.update({
                'partner_ids': [Command.set(partners.ids)],
                'rule_ids': [Command.create(dict(rule['vals'], sequence=seq,
                                                 tier_ids=[Command.create(t) for t in rule['tiers']]))
                             for seq, rule in enumerate(agr['rules'].values())],
            })
            created |= Agreement.create(vals)
        lines_before = len(created.line_ids)
        if self.approve:
            created.action_submit()
            created.action_approve()
        self.result = 'Đã nhập %d thỏa thuận: %s.\nDòng lũy kế dựng lại: %d.' % (
            len(created), ', '.join(created.mapped('name')), len(created.line_ids) - lines_before)
        return {
            'type': 'ir.actions.act_window',
            'name': 'Thỏa thuận đã nhập',
            'res_model': 'rebate.agreement',
            'view_mode': 'list,form',
            'domain': [('id', 'in', created.ids)],
        }
